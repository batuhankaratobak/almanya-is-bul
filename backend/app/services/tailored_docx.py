"""Generate a one-page A4 Word CV tailored automatically to one job posting."""

from __future__ import annotations

import io
from typing import Any, Dict, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.cv_llm import rewrite_tailored_cv
from app.services.cv_profile import load_cv_profile
from app.services.cv_tailor import tailor_cv_for_job
from app.services.cv_translate import detect_output_language, resolve_localized_cv
from app.sources.normalized_job import NormalizedJob


_LABELS = {
    "en": {
        "name_placeholder": "[Your Full Name]",
        "profile": "Profile",
        "fit": "Strengths for this role",
        "skills": "Skills",
        "other": "Also: ",
        "experience": "Experience",
        "projects": "Projects",
        "project_fallback": "Project",
        "education": "Education",
        "languages": "Languages",
        "motivation_heading": "Motivation",
        "motivation": (
            "I am applying for {title}{company_bit}.{skill_bit} "
            "I am looking for long-term growth in Germany."
        ),
        "motivation_skills": " I would especially like to contribute with {skills}.",
        "motivation_fallback": " I want to contribute quickly in an entry-level / junior role.",
        "company_bit": " at {company}",
    },
    "de": {
        "name_placeholder": "[Vor- und Nachname]",
        "profile": "Profil",
        "fit": "Stärken für diese Stelle",
        "skills": "Kenntnisse",
        "other": "Weitere: ",
        "experience": "Erfahrung",
        "projects": "Projekte",
        "project_fallback": "Projekt",
        "education": "Ausbildung",
        "languages": "Sprachen",
        "motivation_heading": "Motivation",
        "motivation": (
            "Ich bewerbe mich auf {title}{company_bit}.{skill_bit} "
            "Mein Ziel ist eine langfristige Entwicklung in Deutschland."
        ),
        "motivation_skills": " Besonders einbringen möchte ich {skills}.",
        "motivation_fallback": (
            " In einer Entry-Level-/Junior-Rolle möchte ich schnell Beitrag leisten."
        ),
        "company_bit": " bei {company}",
    },
}


def build_tailored_docx_bytes(
    job: NormalizedJob | Any,
    *,
    cv_profile: Dict[str, Any] | None = None,
    lang: Optional[str] = None,
) -> Tuple[bytes, str]:
    """Return (docx_bytes, output_lang) where output_lang is 'en' or 'de'."""
    raw_profile = cv_profile or load_cv_profile()
    title = getattr(job, "title", None) or ""
    company = getattr(job, "company", None) or ""
    description = getattr(job, "description", None) or ""
    job_language = getattr(job, "job_language", None)

    output_lang = detect_output_language(
        job_language=job_language,
        title=title,
        description=description,
        override=lang,
    )
    labels = _LABELS[output_lang]
    profile = resolve_localized_cv(raw_profile, output_lang)
    tailored = tailor_cv_for_job(
        profile,
        title=title,
        description=description,
        lang=output_lang,
    )
    tailored = rewrite_tailored_cv(
        tailored,
        title=title,
        company=company,
        description=description,
        lang=output_lang,
    )

    doc = Document()
    _configure_a4(doc)

    name = profile.get("full_name") or labels["name_placeholder"]
    contact_bits = [
        profile.get("city") or "",
        profile.get("email") or "",
        profile.get("phone") or "",
        profile.get("linkedin_or_github") or "",
    ]
    contact = " · ".join(bit for bit in contact_bits if bit)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(name)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    _tighten(h, after=1)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(tailored.headline)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    _tighten(sub, after=0)

    if contact:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = c.add_run(contact)
        cr.font.size = Pt(8)
        cr.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        _tighten(c, after=2)
        _bottom_border(c)

    _heading(doc, labels["profile"])
    _body(doc, tailored.summary)

    # Keep fit section short: max 3 lines, no system/meta wording in the CV body.
    if tailored.requirement_lines:
        _heading(doc, labels["fit"])
        for line in tailored.requirement_lines[:3]:
            _bullet(doc, line)

    _heading(doc, labels["skills"])
    skill_line = ", ".join(tailored.skills[:12]) if tailored.skills else ", ".join(tailored.other_skills[:10])
    _body(doc, skill_line or "—")

    if tailored.experiences:
        _heading(doc, labels["experience"])
        for exp in tailored.experiences:
            label = " — ".join(
                part
                for part in [
                    str(exp.get("title") or "").strip(),
                    str(exp.get("org") or "").strip(),
                    str(exp.get("period") or "").strip(),
                ]
                if part
            )
            if label:
                p = doc.add_paragraph()
                r = p.add_run(label)
                r.bold = True
                r.font.size = Pt(9)
                _tighten(p, before=2, after=0)
            for bullet in list(exp.get("bullets") or [])[:4]:
                _bullet(doc, str(bullet))

    if tailored.projects:
        _heading(doc, labels["projects"])
        for project in tailored.projects:
            name_p = str(project.get("name") or labels["project_fallback"]).strip()
            stack = project.get("stack") if isinstance(project.get("stack"), list) else []
            label = name_p
            if stack:
                label += f" ({', '.join(str(s) for s in stack[:5])})"
            p = doc.add_paragraph()
            r = p.add_run(label)
            r.bold = True
            r.font.size = Pt(9)
            _tighten(p, before=2, after=0)
            for bullet in list(project.get("bullets") or [])[:3]:
                _bullet(doc, str(bullet))

    education = profile.get("education") or []
    if education:
        _heading(doc, labels["education"])
        for edu in education[:2]:
            line = " — ".join(
                part
                for part in [
                    str(edu.get("degree") or "").strip(),
                    str(edu.get("school") or "").strip(),
                    str(edu.get("period") or "").strip(),
                ]
                if part
            )
            if line:
                _body(doc, line)
            details = str(edu.get("details") or "").strip()
            if details:
                _body(doc, details, italic=True)

    languages = profile.get("languages") or []
    if languages:
        _heading(doc, labels["languages"])
        lang_bits = []
        for lang_item in languages:
            n = str(lang_item.get("name") or "").strip()
            level = str(lang_item.get("level") or "").strip()
            if n and level:
                lang_bits.append(f"{n} ({level})")
            elif n:
                lang_bits.append(n)
        _body(doc, " · ".join(lang_bits))

    _heading(doc, labels["motivation_heading"])
    _body(
        doc,
        _motivation(
            title=title,
            company=company,
            skills=tailored.skills,
            labels=labels,
        ),
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue(), output_lang


def _configure_a4(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(12)
    section.right_margin = Mm(12)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def _tighten(paragraph, *, before: float = 0, after: float = 0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    _tighten(p, before=5, after=1)
    _bottom_border(p)


def _body(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.italic = italic
    _tighten(p, after=1)


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    _tighten(p, after=0)
    p.paragraph_format.left_indent = Mm(3)
    p.paragraph_format.line_spacing = 1.0


def _bottom_border(paragraph) -> None:
    """Underline a paragraph without adding an extra empty spacer line."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CBD5E1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _motivation(
    *,
    title: str,
    company: str,
    skills: list,
    labels: Dict[str, str],
) -> str:
    company_bit = labels["company_bit"].format(company=company) if company else ""
    if skills:
        skill_bit = labels["motivation_skills"].format(skills=", ".join(list(skills)[:4]))
    else:
        skill_bit = labels["motivation_fallback"]
    return labels["motivation"].format(
        title=title or "this role",
        company_bit=company_bit,
        skill_bit=skill_bit,
    )
