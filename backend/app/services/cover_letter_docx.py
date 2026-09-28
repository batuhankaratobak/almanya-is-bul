"""Generate a short A4 Anschreiben / cover letter tailored to one job."""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.cv_llm import rewrite_cover_paragraphs
from app.services.cv_profile import load_cv_profile
from app.services.cv_tailor import tailor_cv_for_job
from app.services.cv_translate import detect_output_language, resolve_localized_cv
from app.sources.normalized_job import NormalizedJob


def build_cover_letter_docx_bytes(
    job: NormalizedJob | Any,
    *,
    cv_profile: Dict[str, Any] | None = None,
    lang: Optional[str] = None,
) -> Tuple[bytes, str]:
    raw = cv_profile or load_cv_profile()
    title = getattr(job, "title", None) or ""
    company = getattr(job, "company", None) or ""
    location = getattr(job, "location", None) or ""
    description = getattr(job, "description", None) or ""
    job_language = getattr(job, "job_language", None)

    output_lang = detect_output_language(
        job_language=job_language,
        title=title,
        description=description,
        override=lang,
    )
    profile = resolve_localized_cv(raw, output_lang)
    tailored = tailor_cv_for_job(
        profile,
        title=title,
        description=description,
        lang=output_lang,
    )
    focus = tailored.focus
    skills = tailored.skills[:4]

    name = profile.get("full_name") or "Your Name"
    email = profile.get("email") or ""
    phone = profile.get("phone") or ""
    city = profile.get("city") or ""
    today = date.today().strftime("%d.%m.%Y")

    doc = Document()
    _configure_a4(doc)

    # Sender block
    for line in [name, city, phone, email]:
        if not line:
            continue
        p = doc.add_paragraph()
        r = p.add_run(str(line))
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)

    # Recipient
    if company:
        p = doc.add_paragraph()
        r = p.add_run(company)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)
    if location:
        p = doc.add_paragraph()
        r = p.add_run(location)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)

    # Date (right-aligned in DE tradition often left is fine for modern)
    p = doc.add_paragraph()
    r = p.add_run(f"{city.split('·')[0].strip() + ', ' if city else ''}{today}".strip(", "))
    r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(12)

    # Subject
    subject = _subject(output_lang, title)
    p = doc.add_paragraph()
    r = p.add_run(subject)
    r.bold = True
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(12)

    # Salutation + body
    fallback = _body_paragraphs(
        lang=output_lang,
        title=title,
        company=company,
        focus=focus,
        skills=skills,
    )
    paragraphs = rewrite_cover_paragraphs(
        title=title,
        company=company,
        description=description,
        lang=output_lang,
        skills=skills,
        experiences=tailored.experiences,
        projects=tailored.projects,
        fallback=fallback,
    )
    for paragraph in paragraphs:
        p = doc.add_paragraph()
        r = p.add_run(paragraph)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15

    # Closing
    closing = (
        "Mit freundlichen Grüßen"
        if output_lang == "de"
        else "Kind regards"
    )
    p = doc.add_paragraph()
    r = p.add_run(closing)
    r.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    r = p.add_run(name)
    r.font.size = Pt(10)

    # Attachments note (German standard)
    p = doc.add_paragraph()
    note = (
        "Anlagen: Lebenslauf"
        if output_lang == "de"
        else "Enclosures: Curriculum Vitae"
    )
    r = p.add_run(note)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    p.paragraph_format.space_before = Pt(16)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue(), output_lang


def _configure_a4(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def _subject(lang: str, title: str) -> str:
    role = (title or "").strip() or ("die ausgeschriebene Stelle" if lang == "de" else "the advertised role")
    if lang == "de":
        return f"Bewerbung als {role}"
    return f"Application for {role}"


def _body_paragraphs(
    *,
    lang: str,
    title: str,
    company: str,
    focus: str,
    skills: list,
) -> list[str]:
    role = (title or "").strip()
    company_bit = company.strip() if company else ""
    skill_bit = ", ".join(skills) if skills else ""

    if lang == "de":
        greet = "Sehr geehrte Damen und Herren,"
        p1 = (
            f"hiermit bewerbe ich mich um die Position „{role}“"
            + (f" bei {company_bit}" if company_bit else "")
            + ". Als Absolvent der Computer Engineering / Informatik bringe ich "
            "praktische IT- und Softwareerfahrung sowie eine strukturierte, zuverlässige Arbeitsweise mit."
        )
        angle = {
            "it_support": (
                "In kommunalen IT-Abteilungen und operativen Rollen habe ich Nutzeranfragen, "
                "Windows-/Netzwerkthemen sowie Helpdesk-, ERP- und CRM-Prozesse begleitet."
            ),
            "software": (
                "In eigenen Projekten habe ich End-to-End Software mit Python, React/TypeScript "
                "und SQL-basierten APIs umgesetzt — inklusive Tests und klarer Dokumentation."
            ),
            "mobile": (
                "Ich habe Offline-First-Mobile-Prototypen mit TypeScript und React Native/Expo "
                "gebaut und dabei UX, lokale Datenhaltung und Qualitätssicherung priorisiert."
            ),
            "realtime": (
                "Ich habe Echtzeit-Full-Stack-Systeme mit Node.js und Socket.IO umgesetzt, "
                "inklusive serverseitiger Validierung und robustem Reconnect-Verhalten."
            ),
            "data_research": (
                "Ich habe operative Prozesse mit ERP/CRM sowie Python-Automatisierung und "
                "Excel-/Datenlieferungen verbunden und dabei auf Datenqualität geachtet."
            ),
            "vision": (
                "Ich habe Python-Desktop- und Computer-Vision-Pipelines mit Fokus auf "
                "Zuverlässigkeit und nachvollziehbare Testergebnisse entwickelt."
            ),
            "general": (
                "Mein Profil verbindet kommunale IT-Support-Erfahrung mit praxisnahen "
                "Softwareprojekten und einer schnellen Einarbeitung in neue Tools."
            ),
        }.get(focus) or (
            "Mein Profil verbindet IT-Support-Erfahrung mit praxisnahen Softwareprojekten."
        )
        p2 = angle
        if skill_bit:
            p2 += f" Besonders relevant für diese Stelle sind u. a.: {skill_bit}."
        p3 = (
            "Ich möchte mich langfristig in Deutschland beruflich weiterentwickeln und "
            "in einem Junior-/Einstiegsrahmen schnell Verantwortung übernehmen. "
            "Über die Einladung zu einem Gespräch freue ich mich sehr."
        )
        return [greet, p1, p2, p3]

    greet = "Dear Hiring Team,"
    p1 = (
        f"I am writing to apply for the {role} position"
        + (f" at {company_bit}" if company_bit else "")
        + ". I am a Computer Engineering graduate with hands-on IT support experience "
        "and practical software projects delivered end to end."
    )
    angle = {
        "it_support": (
            "In municipal IT teams and operational roles I handled user requests, "
            "Windows/network troubleshooting, and Help Desk / ERP / CRM workflows."
        ),
        "software": (
            "I have built end-to-end software with Python, React/TypeScript and SQL-backed APIs, "
            "including testing and clear documentation."
        ),
        "mobile": (
            "I built offline-first mobile prototypes with TypeScript and React Native/Expo, "
            "with attention to UX, local storage and QA."
        ),
        "realtime": (
            "I built real-time full-stack systems with Node.js and Socket.IO, including "
            "server-side validation and reconnection handling."
        ),
        "data_research": (
            "I combined operational ERP/CRM workflows with Python automation and "
            "Excel/data delivery, focusing on data quality."
        ),
        "vision": (
            "I developed Python desktop and computer-vision pipelines with an emphasis "
            "on reliability and testable results."
        ),
        "general": (
            "My profile combines municipal IT support experience with practical software projects "
            "and fast ramp-up on new tools."
        ),
    }.get(focus) or (
        "My profile combines IT support experience with practical software projects."
    )
    p2 = angle
    if skill_bit:
        p2 += f" Skills especially relevant to this role include: {skill_bit}."
    p3 = (
        "I am looking for long-term growth in Germany and want to contribute quickly "
        "in a junior / entry-level role. I would welcome the opportunity to discuss the position."
    )
    return [greet, p1, p2, p3]
